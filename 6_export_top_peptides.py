#-----------------------------------------------------------------------------
# 6_export_top_peptides.py
#-----------------------------------------------------------------------------
# Reads merged_counts.csv, filters out stop-codon peptides, ranks by final-round
# count (e.g. R8), exports top N peptides to CSV and Excel.
# Streams the file (one pass) and keeps only top N in a heap — no full load.
# Run after 4_merge_rounds.py.
#-----------------------------------------------------------------------------
# Input:   data/output/merged_counts.csv, config.json
# Output:  data/output/top100_peptides.csv, data/output/top100_peptides.xlsx
#-----------------------------------------------------------------------------

import os
import json
import csv
import heapq

_script_dir = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(_script_dir, "config.json")) as f:
    cfg = json.load(f)

OUTPUT_DIR = cfg["output_dir"]
ROUNDS = cfg["rounds"]
TOP_N = cfg.get("top_n_peptides", 100)

merged_path = os.path.join(OUTPUT_DIR, "merged_counts.csv")
out_csv = os.path.join(OUTPUT_DIR, f"top{TOP_N}_peptides.csv")
out_xlsx = os.path.join(OUTPUT_DIR, f"top{TOP_N}_peptides.xlsx")

if not os.path.isfile(merged_path):
    print(f"Error: {merged_path} not found. Run 4_merge_rounds.py first.")
    exit(1)

# Build output column names from config rounds
count_cols = [f"count_r{r}" for r in ROUNDS]
freq_cols = [f"freq_r{r}" for r in ROUNDS]
last_round = ROUNDS[-1]
count_last = f"count_r{last_round}"

def get_count(row, r):
    try:
        return int(row.get(f"count_r{r}", 0) or 0)
    except (ValueError, TypeError):
        return 0

# Stream merged_counts.csv: keep only top N by final-round count (min-heap of size N)
# Heap stores (count_r8, tie_breaker, row); tie_breaker avoids comparing dicts when counts equal.
top_heap = []
n_seen = 0
with open(merged_path) as f:
    reader = csv.DictReader(f)
    for row in reader:
        pep = (row.get("peptide") or "").strip()
        if "*" in pep:
            continue
        n_seen += 1
        c_last = get_count(row, last_round)
        if len(top_heap) < TOP_N:
            heapq.heappush(top_heap, (c_last, n_seen, row))
        elif c_last > top_heap[0][0]:
            heapq.heapreplace(top_heap, (c_last, n_seen, row))

# Sort descending for output (top count first)
top = [row for (_, _, row) in sorted(top_heap, key=lambda x: -x[0])]
print(f"Scanned {n_seen} peptides (no stop), kept top {len(top)} by R{last_round} count.")

# Output column order: peptide, R1_count, R5_count, ..., freq_R1, ..., enrichment, rank
out_count_names = [f"R{r}_count" for r in ROUNDS]
out_freq_names = [f"freq_R{r}" for r in ROUNDS]
out_header = ["peptide"] + out_count_names + out_freq_names + ["enrichment_R8_R1", f"rank_R{last_round}"]

# Write CSV
with open(out_csv, "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(out_header)
    for rank, row in enumerate(top, 1):
        pep = (row.get("peptide") or "").strip()
        counts = [str(row.get(c, "")) for c in count_cols]
        freqs = [str(row.get(c, "")) for c in freq_cols]
        enr = (row.get("enrichment_r8_vs_r1") or "").strip()
        w.writerow([pep] + counts + freqs + [enr, rank])
print(f"Wrote {out_csv}")

# Write Excel if openpyxl available
try:
    import openpyxl
    from openpyxl.styles import Font, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Top peptides"
    for c, h in enumerate(out_header, 1):
        ws.cell(row=1, column=c, value=h)
        ws.cell(row=1, column=c).font = Font(bold=True)
    for r, row in enumerate(top, 2):
        pep = (row.get("peptide") or "").strip()
        counts = [row.get(c, "") for c in count_cols]
        freqs = [row.get(c, "") for c in freq_cols]
        enr = (row.get("enrichment_r8_vs_r1") or "").strip()
        for c, val in enumerate([pep] + counts + freqs + [enr, r - 1], 1):
            ws.cell(row=r, column=c, value=val)
    wb.save(out_xlsx)
    print(f"Wrote {out_xlsx}")
except ImportError:
    print("openpyxl not installed; skipping Excel. Install with: pip install openpyxl")