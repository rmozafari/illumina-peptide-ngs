#-----------------------------------------------------------------------------
# 7_conserved_positions.py
#-----------------------------------------------------------------------------
# Reads top N peptides CSV, extracts the 10 variable positions (between the
# two C's in MC...C), and reports per-position amino acid counts to detect
# conserved positions. Run after 6_export_top_peptides.py.
#-----------------------------------------------------------------------------
# Input:   data/output/top{N}_peptides.csv, config.json
# Output:  data/output/conserved_positions.csv, conserved_positions.txt, conserved_positions.xlsx (if openpyxl)
#-----------------------------------------------------------------------------

import os
import json
import csv
from collections import Counter

_script_dir = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(_script_dir, "config.json")) as f:
    cfg = json.load(f)

OUTPUT_DIR = cfg["output_dir"]
TOP_N = cfg.get("top_n_peptides", 1000)

# Peptide is 13 aa: M C X10 C → variable positions are indices 2..11 (0-based)
VAR_START = 2
VAR_LEN = 10

in_csv = os.path.join(OUTPUT_DIR, f"top{TOP_N}_peptides.csv")
out_csv = os.path.join(OUTPUT_DIR, "conserved_positions.csv")
out_txt = os.path.join(OUTPUT_DIR, "conserved_positions.txt")
out_xlsx = os.path.join(OUTPUT_DIR, "conserved_positions.xlsx")

if not os.path.isfile(in_csv):
    print(f"Error: {in_csv} not found. Run 6_export_top_peptides.py first (with top_n_peptides={TOP_N}).")
    exit(1)

peptides = []
with open(in_csv) as f:
    reader = csv.DictReader(f)
    for row in reader:
        pep = (row.get("peptide") or "").strip()
        if len(pep) >= VAR_START + VAR_LEN:
            peptides.append(pep)

if not peptides:
    print("No peptides with 10 variable positions found.")
    exit(1)

# Per-position counts (position 0 = first variable, i.e. 3rd aa in full peptide)
position_counts = [Counter() for _ in range(VAR_LEN)]
for pep in peptides:
    var = pep[VAR_START : VAR_START + VAR_LEN]
    for i, aa in enumerate(var):
        if i < VAR_LEN:
            position_counts[i][aa] += 1

n = len(peptides)
rows = []
for pos in range(VAR_LEN):
    c = position_counts[pos]
    most = c.most_common(3)
    top_aa = most[0][0] if most else ""
    top_cnt = most[0][1] if most else 0
    top_pct = 100.0 * top_cnt / n if n else 0
    second_aa = most[1][0] if len(most) > 1 else ""
    second_cnt = most[1][1] if len(most) > 1 else 0
    second_pct = 100.0 * second_cnt / n if n else 0
    third_aa = most[2][0] if len(most) > 2 else ""
    third_pct = 100.0 * most[2][1] / n if n and len(most) > 2 else 0
    # Simple conservation: top residue >= 30% we call "conserved"
    conserved = "yes" if top_pct >= 30 else "no"
    rows.append({
        "position": pos + 1,
        "top_residue": top_aa,
        "top_pct": round(top_pct, 1),
        "second_residue": second_aa,
        "second_pct": round(second_pct, 1),
        "third_residue": third_aa,
        "third_pct": round(third_pct, 1),
        "conserved": conserved,
   })

with open(out_csv, "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=["position", "top_residue", "top_pct", "second_residue", "second_pct", "third_residue", "third_pct", "conserved"])
    w.writeheader()
    w.writerows(rows)
print(f"Wrote {out_csv}")

# Text summary
with open(out_txt, "w") as f:
    f.write(f"Conserved positions in top {len(peptides)} peptides (10 variable positions)\n")
    f.write("=" * 60 + "\n\n")
    consensus = "".join(r["top_residue"] for r in rows)
    f.write(f"Consensus (most frequent residue per position): {consensus}\n\n")
    f.write("Position  Top   %     Second  %     Third   %     Conserved?\n")
    f.write("-" * 60 + "\n")
    for r in rows:
        f.write(f"   {r['position']:2d}       {r['top_residue']}   {r['top_pct']:5.1f}   {r['second_residue']}     {r['second_pct']:5.1f}   {r['third_residue']}     {r['third_pct']:5.1f}   {r['conserved']}\n")
    f.write("\n(Conserved = top residue >= 30%)\n")
print(f"Wrote {out_txt}")

# Write Excel if openpyxl available
fieldnames = ["position", "top_residue", "top_pct", "second_residue", "second_pct", "third_residue", "third_pct", "conserved"]
try:
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Conserved positions"
    for c, h in enumerate(fieldnames, 1):
        ws.cell(row=1, column=c, value=h)
        ws.cell(row=1, column=c).font = Font(bold=True)
    for r, row in enumerate(rows, 2):
        for c, key in enumerate(fieldnames, 1):
            ws.cell(row=r, column=c, value=row.get(key, ""))
    wb.save(out_xlsx)
    print(f"Wrote {out_xlsx}")
except ImportError:
    print("openpyxl not installed; skipping Excel. Install with: pip install openpyxl")

# Print consensus to console
print(f"Top {len(peptides)} peptides: consensus (10 variable positions) = {consensus}")
conserved_pos = [r["position"] for r in rows if r["conserved"] == "yes"]
if conserved_pos:
    print(f"Conserved positions (top >= 30%): {conserved_pos}")
else:
    print("No position with a single residue >= 30%.")
